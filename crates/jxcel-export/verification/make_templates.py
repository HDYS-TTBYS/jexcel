#!/usr/bin/env python3
"""確認用のテンプレート（verification/templates/*.xlsx）を作り直す。

使い方: python3 make_templates.py 出力先ディレクトリ [--uno ポート]
  openpyxl が要る。--uno を付けると、LibreOffice で本物のピボットテーブルも作る。その LibreOffice は、変換用の
  `soffice --convert-to` と取り合わないよう、別のプロファイルで起動しておく:
    soffice -env:UserInstallation=file:///tmp/lo-uno --headless --invisible --norestore \\
            --accept="socket,host=localhost,port=2007;urp;" &
  チャートのキャッシュ付きのテンプレートも LibreOffice の変換を使う（soffice が PATH にあること）。
"""
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumDataSource
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

LOOP_A = "{{#each 明細}}{{品目}}"
LOOP_B = "{{数}}"


def loop_sheet(ws):
    ws["A1"] = "見出し"; ws["B1"] = 100
    ws["A2"] = LOOP_A; ws["B2"] = LOOP_B
    ws["A3"] = "末尾"; ws["B3"] = 7


def three_d(out):
    wb = openpyxl.Workbook()
    a = wb.active; a.title = "A"
    for r in range(1, 6): a.cell(r, 2, r * 10)
    m = wb.create_sheet("明細"); loop_sheet(m)
    c = wb.create_sheet("C")
    for r in range(1, 6): c.cell(r, 2, r * 1000)
    s = wb.create_sheet("集計")
    s["A1"] = "=SUM(A:C!B3)"      # ずれ方が食い違う → 集計関数なのでシートごとに分かれる
    s["A2"] = "=MAX(A:C!B1:B3)"
    wb.save(os.path.join(out, "three-d-sum.xlsx"))

    wb = openpyxl.Workbook()
    a = wb.active; a.title = "A"; loop_sheet(a)
    m = wb.create_sheet("明細"); loop_sheet(m)
    c = wb.create_sheet("C")
    for r in range(1, 6): c.cell(r, 2, r * 1000)
    s = wb.create_sheet("集計")
    s["A1"] = "=SUM(A:明細!B3)"   # 範囲内の動いたシートがどれも同じずれ方 → そのままの 3D 参照
    s["A2"] = "=SUM(A:明細!B1:B3)"
    s["A3"] = "=SUM(A:C!B1)"      # ループより前の行 → 動かない
    wb.save(os.path.join(out, "three-d-both-loop-sheets.xlsx"))


def absolute_anchor(out, heights):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "明細"
    ws.sheet_format.defaultRowHeight = 15
    ws["A1"] = "見出し"; ws["B1"] = 100
    ws["A2"] = LOOP_A; ws["B2"] = LOOP_B
    ws["A3"] = "末尾"; ws["B3"] = 7
    if heights:
        ws["A4"] = "次"; ws["A5"] = "隠れ"; ws["A6"] = "最後"; ws["B6"] = 9
        for r, h in {1: 30, 2: 20, 3: 40, 6: 25}.items(): ws.row_dimensions[r].height = h
        ws.row_dimensions[5].hidden = True
    P = 12700; E = 190500

    def chart(x, y, cx, cy):
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3)); ch.legend = None
        ch.anchor = AbsoluteAnchor(pos=XDRPoint2D(x, y), ext=XDRPositiveSize2D(cx, cy)); ws.add_chart(ch)

    if heights:
        chart(1800000, 60 * P, 1200000, 20 * P)    # 末尾の行（3 行目）の中
        chart(3200000, 40 * P, 1200000, 100 * P)   # 2 行目の中ほどから、最後の行より後ろまで
        ws.title = "明細"
        wb.save(os.path.join(out, "absolute-anchor-row-heights.xlsx"))
    else:
        chart(1800000, 2 * E, 1500000, 2 * E)      # 末尾の行の隣
        chart(3400000, 0, 1500000, 3 * E)          # 1〜3 行目に掛かる
        wb.save(os.path.join(out, "absolute-anchor.xlsx"))


X14 = '''<extLst>
<ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:conditionalFormattings><x14:conditionalFormatting xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:cfRule type="cellIs" priority="1" operator="greaterThan"><xm:f>明細!$B$3</xm:f></x14:cfRule><xm:sqref>B1:B3</xm:sqref></x14:conditionalFormatting></x14:conditionalFormattings></ext>
<ext uri="{CCE6A557-97BC-4b89-ADB6-D93E4EDE0EAF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:dataValidations xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main" count="1"><x14:dataValidation type="list" allowBlank="1" showInputMessage="1" showErrorMessage="1"><x14:formula1><xm:f>明細!$A$1:$A$3</xm:f></x14:formula1><xm:sqref>E1</xm:sqref></x14:dataValidation></x14:dataValidations></ext>
<ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:sparklineGroup type="column"><x14:colorSeries rgb="FF376092"/><x14:sparklines><x14:sparkline><xm:f>明細!B1:B3</xm:f><xm:sqref>D1</xm:sqref></x14:sparkline><x14:sparkline><xm:f>明細!B3</xm:f><xm:sqref>D3</xm:sqref></x14:sparkline></x14:sparklines></x14:sparklineGroup></x14:sparklineGroups></ext>
</extLst>'''


def x14(out):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "明細"; loop_sheet(ws)
    tmp = os.path.join(tempfile.mkdtemp(), "x14_0.xlsx"); wb.save(tmp)
    zi = zipfile.ZipFile(tmp); zo = zipfile.ZipFile(os.path.join(out, "x14-sparkline.xlsx"), "w", zipfile.ZIP_DEFLATED)
    for i in zi.infolist():
        b = zi.read(i.filename)
        if i.filename == "xl/worksheets/sheet1.xml":
            b = b.decode().replace("</worksheet>", X14 + "</worksheet>").encode()
        zo.writestr(i, b)
    zo.close()


def chart_cf_dv(out):
    """グラフ・条件付き書式・入力規則・定義名・印刷範囲。キャッシュを持たせるため LibreOffice で保存し直す。"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "明細"
    ws.append(["品目", "数量"]); ws.append(["先頭", 1]); ws.append([LOOP_A, LOOP_B]); ws.append(["末尾", 5]); ws.append(["合計", "=SUM(B2:B4)"])
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=2, min_row=2, max_row=4)); ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(ch, "D2")
    ws.conditional_formatting.add("B2:B4", CellIsRule(operator="greaterThan", formula=["2"], fill=PatternFill("solid", bgColor="FFFF00")))
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="B5"); ws.add_data_validation(dv); dv.add("B2:B4")
    wb.defined_names["範囲"] = DefinedName("範囲", attr_text="明細!$A$2:$B$4")
    ws.print_area = "A1:B5"
    d = tempfile.mkdtemp(); src = os.path.join(d, "chart-conditional-format-validation.xlsx"); wb.save(src)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(d, "lo"), src], check=True, capture_output=True)
    shutil.copy(os.path.join(d, "lo", os.path.basename(src)), out)


def pivots(out, port):
    """LibreOffice（UNO）で本物のピボットテーブルを作る。"""
    import uno
    from com.sun.star.beans import PropertyValue
    from com.sun.star.sheet.DataPilotFieldOrientation import ROW, DATA
    from com.sun.star.sheet.GeneralFunction import SUM
    from com.sun.star.table import CellAddress, CellRangeAddress

    ctx = uno.getComponentContext()
    r = ctx.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", ctx)
    c = r.resolve(f"uno:socket,host=localhost,port={port};urp;StarOffice.ComponentContext")
    desk = c.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", c)

    def build(name, separate_sheet):
        pv = PropertyValue(); pv.Name = "Hidden"; pv.Value = True
        doc = desk.loadComponentFromURL("private:factory/scalc", "_blank", 0, (pv,))
        sh = doc.Sheets; d = sh.getByIndex(0); d.Name = "データ"
        for i, (a, b) in enumerate([("品目", "数量"), ("先頭", 1), (LOOP_A, LOOP_B), ("末尾", 5)]):
            d.getCellByPosition(0, i).String = a
            cb = d.getCellByPosition(1, i)
            cb.Value = b if isinstance(b, int) else 0
            if not isinstance(b, int): cb.String = b
        a = CellRangeAddress(); a.Sheet = 0; a.StartColumn = 0; a.StartRow = 0; a.EndColumn = 1; a.EndRow = 3
        if separate_sheet:
            sh.insertNewByName("集計", 1); target = sh.getByName("集計"); dst_row = 0
        else:
            d.getCellByPosition(0, 4).String = "合計"; d.getCellByPosition(1, 4).Formula = "=SUM(B2:B4)"
            target = d; dst_row = 7
        dp = target.DataPilotTables
        desc = dp.createDataPilotDescriptor(); desc.SourceRange = a
        desc.getDataPilotFields().getByIndex(0).Orientation = ROW
        f = desc.getDataPilotFields().getByIndex(1); f.Orientation = DATA; f.Function = SUM
        dst = CellAddress(); dst.Sheet = target.RangeAddress.Sheet; dst.Column = 0; dst.Row = dst_row
        dp.insertNewByName("pv", dst, desc)
        p = PropertyValue(); p.Name = "FilterName"; p.Value = "Calc MS Excel 2007 XML"
        doc.storeToURL("file://" + os.path.abspath(os.path.join(out, name)), (p,))
        doc.close(True)

    build("pivot.xlsx", True)
    build("pivot-in-loop-sheet.xlsx", False)


def header_footer_docx(out):
    """ヘッダー・フッターの段落のループ。"""
    import docx

    d = docx.Document()
    d.add_paragraph("本文 {{請求番号}}")
    sec = d.sections[0]
    h = sec.header
    h.paragraphs[0].text = "{{#each 明細}}ヘッダー {{品目}}"
    h.add_paragraph("{{/each}}")
    f = sec.footer
    f.paragraphs[0].text = "{{#each 明細}}フッター {{_n}}: {{品目}}"
    f.add_paragraph("{{/each}}")
    d.save(os.path.join(out, "header-footer.docx"))


if __name__ == "__main__":
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    header_footer_docx(out); three_d(out); absolute_anchor(out, False); absolute_anchor(out, True); x14(out); chart_cf_dv(out)
    if "--uno" in sys.argv:
        pivots(out, sys.argv[sys.argv.index("--uno") + 1])
